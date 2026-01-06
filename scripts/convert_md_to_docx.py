from docx import Document
from pathlib import Path

src = Path('docs/Documentation.md')
dst = Path('docs/Documentation.docx')

doc = Document()

for line in src.read_text(encoding='utf-8').splitlines():
    if line.startswith('# '):
        doc.add_heading(line.lstrip('# ').strip(), level=1)
    elif line.startswith('## '):
        doc.add_heading(line.lstrip('# ').strip(), level=2)
    elif line.startswith('- '):
        doc.add_paragraph(line.lstrip('- ').strip(), style='List Bullet')
    elif line.strip() == '':
        doc.add_paragraph('')
    else:
        doc.add_paragraph(line)

doc.save(dst)
print('Wrote', dst)
